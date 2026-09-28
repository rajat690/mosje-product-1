"""User dataset uploads: validate a CBSE or Jan Aadhaar file and replace that table's contents.

Accepted files: .csv (UTF-8, or Windows-1252 as a fallback), .csv.gz and .xlsx (first sheet).
Column names are the ones of the bundled synthetic files (data/*.csv.gz); matching ignores case,
spaces and underscores, so "Roll No" is read as roll_no and "nameeng" as nameEng.

Memory: the file is never loaded as a whole. It is read row by row twice (1: validate, 2: COPY into
Postgres inside one transaction), so a few hundred thousand rows fit easily in 512 MB. The only thing
kept in memory is the set of key values used to detect duplicates.
"""
from __future__ import annotations

import csv
import gzip
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from typing import IO, Iterator, Optional

from sqlalchemy import text
from sqlalchemy.engine import Engine

from mosje.eligibility import derive_education_stage
from mosje.normalize import l1_dob, l1_gender

from .bulk import bulk_insert

EXAM_YEAR = "2025-26"
MAX_BAD_ROW_SAMPLE = 10
RESULT_TABLES = ("eligibility_results", "student_eligibility", "outreach_queue", "link_decisions")
CATEGORY_CANON = {"sc": "SC", "st": "ST", "obc": "OBC", "general": "General", "gen": "General",
                  "minority": "Minority"}
GENDER_TITLE = {"male": "Male", "female": "Female", "transgender": "Transgender", "other": "Other"}


class UploadError(ValueError):
    """The file cannot be read at all (wrong type, empty, not a table)."""


class UploadBusy(RuntimeError):
    """A pipeline run is in progress; uploads would change its inputs."""


@dataclass(frozen=True)
class Spec:
    dataset: str                 # table name
    label: str
    key: str                     # unique column
    must_fill: tuple             # header required AND every row needs a value
    required: tuple              # header required, value may be blank
    optional: tuple              # header may be left out, value may be blank

    @property
    def columns(self) -> tuple:
        return self.must_fill + self.required + self.optional


CBSE = Spec("cbse_results", "CBSE passed students", "roll_no",
            must_fill=("roll_no", "candidate_name", "dob", "gender", "class_passed"),
            required=("father_name", "mother_name"),
            optional=("apaar_id", "school", "district", "mobile"))
JAN_AADHAAR = Spec("jan_aadhaar_members", "Jan Aadhaar members", "member_id",
                   must_fill=("jan_aadhaar_id", "member_id", "nameEng", "dob", "gender"),
                   required=("fatherNameEng", "motherNameEng", "category", "annual_family_income", "domicile_state"),
                   optional=("nameHnd", "district", "disability", "mobile", "relation"))
SPECS = {"cbse": CBSE, "jan-aadhaar": JAN_AADHAAR}
# column order of the bundled synthetic files (used for templates / docs)
FILE_ORDER = {
    "cbse_results": ("roll_no", "apaar_id", "candidate_name", "dob", "gender", "father_name", "mother_name",
                     "class_passed", "school", "district", "mobile"),
    "jan_aadhaar_members": ("jan_aadhaar_id", "member_id", "nameEng", "nameHnd", "dob", "gender", "fatherNameEng",
                            "motherNameEng", "category", "annual_family_income", "district", "domicile_state",
                            "disability", "mobile", "relation"),
}


def _canon(name) -> str:
    return re.sub(r"[^a-z0-9]", "", str(name or "").replace("\ufeff", "").lower())


# ------------------------------------------------------------------ reading
def _cell(v) -> str:
    """Excel/CSV cell -> clean string. 1234.0 -> '1234', dates -> YYYY-MM-DD."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v).strip()


def detect_format(fileobj: IO[bytes], filename: str) -> str:
    fileobj.seek(0)
    head = fileobj.read(4)
    fileobj.seek(0)
    name = (filename or "").lower()
    if not head:
        raise UploadError("the file is empty")
    if head[:2] == b"\x1f\x8b":
        return "csv.gz"
    if head[:2] == b"PK":
        if name.endswith((".xlsx", ".xlsm")) or not name.endswith(".csv"):
            return "xlsx"
        raise UploadError("this looks like a zip file, not a CSV. Upload the .csv itself (or .csv.gz / .xlsx)")
    if head[:4] == b"\xd0\xcf\x11\xe0":
        raise UploadError("old Excel .xls files are not supported. In Excel use File > Save As > "
                          "'CSV UTF-8 (Comma delimited) (*.csv)' or 'Excel Workbook (*.xlsx)'")
    return "csv"


def _csv_rows(fileobj: IO[bytes], fmt: str, encoding: str) -> Iterator[list[str]]:
    fileobj.seek(0)
    raw = gzip.GzipFile(fileobj=fileobj, mode="rb") if fmt == "csv.gz" else fileobj
    txt = io.TextIOWrapper(raw, encoding=encoding, newline="")
    try:
        for row in csv.reader(txt):
            yield [c.strip() for c in row]
    finally:
        txt.detach()


def _xlsx_rows(fileobj: IO[bytes]) -> Iterator[list[str]]:
    from openpyxl import load_workbook
    fileobj.seek(0)
    try:
        wb = load_workbook(fileobj, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise UploadError(f"could not open the Excel file: {e}")
    try:
        ws = wb.worksheets[0]
        for row in ws.iter_rows(values_only=True):
            yield [_cell(v) for v in row]
    finally:
        wb.close()


def iter_rows(fileobj: IO[bytes], fmt: str, encoding: str = "utf-8-sig") -> Iterator[list[str]]:
    return _xlsx_rows(fileobj) if fmt == "xlsx" else _csv_rows(fileobj, fmt, encoding)


def sha256_of(fileobj: IO[bytes]) -> str:
    fileobj.seek(0)
    h = hashlib.sha256()
    for chunk in iter(lambda: fileobj.read(1 << 20), b""):
        h.update(chunk)
    fileobj.seek(0)
    return h.hexdigest()


# ------------------------------------------------------------------ validation
@dataclass
class Report:
    dataset: str
    label: str
    filename: str
    file_format: str = ""
    encoding: str = ""
    rows_read: int = 0
    columns_found: list = field(default_factory=list)
    missing_columns: list = field(default_factory=list)
    extra_columns: list = field(default_factory=list)
    duplicate_columns: list = field(default_factory=list)
    bad_row_count: int = 0
    problem_counts: dict = field(default_factory=dict)
    bad_rows_sample: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    too_many_rows: Optional[int] = None

    @property
    def ok(self) -> bool:
        return not (self.missing_columns or self.duplicate_columns or self.bad_row_count or self.too_many_rows) \
            and self.rows_read > 0

    def as_dict(self) -> dict:
        errs = []
        if self.missing_columns:
            errs.append(f"missing required column(s): {', '.join(self.missing_columns)}")
        if self.duplicate_columns:
            errs.append(f"column(s) appear more than once: {', '.join(self.duplicate_columns)}")
        if self.bad_row_count:
            errs.append(f"{self.bad_row_count:,} row(s) have problems (first {len(self.bad_rows_sample)} shown in "
                        f"bad_rows_sample; 'row' is the line number in the file, header = line 1)")
        if self.too_many_rows:
            errs.append(f"the file has more than {self.too_many_rows:,} data rows (MAX_UPLOAD_ROWS)")
        if not self.rows_read and not self.missing_columns:
            errs.append("the file has a header but no data rows")
        return {"ok": self.ok, "dataset": self.dataset, "label": self.label, "filename": self.filename,
                "file_format": self.file_format, "encoding": self.encoding, "rows_read": self.rows_read,
                "errors": errs, "missing_columns": self.missing_columns, "extra_columns": self.extra_columns,
                "duplicate_columns": self.duplicate_columns, "bad_row_count": self.bad_row_count,
                "problem_counts": self.problem_counts, "bad_rows_sample": self.bad_rows_sample,
                "warnings": self.warnings, "columns_found": self.columns_found}


def _map_header(spec: Spec, header: list[str], rep: Report) -> dict[str, int]:
    wanted = {_canon(c): c for c in spec.columns}
    idx, seen = {}, {}
    for i, h in enumerate(header):
        c = _canon(h)
        if not c:
            continue
        rep.columns_found.append(h)
        if c in seen:
            rep.duplicate_columns.append(h)
            continue
        seen[c] = i
        if c in wanted:
            idx[wanted[c]] = i
        else:
            rep.extra_columns.append(h)
    rep.missing_columns = [c for c in spec.must_fill + spec.required if c not in idx]
    return idx


def _income(v: str) -> tuple[Optional[str], bool]:
    s = re.sub(r"(?i)^(rs\.?|inr|₹)", "", v.replace(",", "").replace(" ", "")).strip()
    if s == "":
        return None, True
    try:
        f = float(s)
    except ValueError:
        return None, False
    if f < 0:
        return None, False
    return str(int(f)), True


def clean_row(spec: Spec, vals: dict[str, str], today: date) -> tuple[dict, list[str], list[str]]:
    """Returns (clean values, errors, warnings) for one row."""
    errors, warns = [], []
    out = dict(vals)
    for c in spec.must_fill:
        if not vals.get(c):
            errors.append(f"{c} is empty")
    dob = vals.get("dob", "")
    if dob:
        iso = l1_dob(dob)
        if not iso:
            errors.append(f"dob '{dob}' is not a date (use e.g. 2009-04-07, 07-04-2009, 07/04/2009 or 07 Apr 2009)")
        elif not (date(1900, 1, 1) <= date.fromisoformat(iso) <= today):
            errors.append(f"dob '{dob}' is outside 1900 to today")
        out["_dob_iso"] = iso
    g = vals.get("gender", "")
    if g:
        gn = l1_gender(g)
        if not gn:
            errors.append(f"gender '{g}' not recognised (use Male/Female/Transgender/Other or M/F/T/O)")
        out["_gender_norm"] = gn
    if spec is CBSE:
        cp = vals.get("class_passed", "")
        if cp:
            stage = derive_education_stage(cp)
            if stage is None:
                errors.append(f"class_passed '{cp}' must be X or XII (10 / 12 also accepted)")
            else:
                out["class_passed"] = "XII" if "Higher Education" in stage else "X"
    else:
        if g and out.get("_gender_norm"):
            out["gender"] = GENDER_TITLE[out["_gender_norm"]]     # eligibility engine expects Male/Female/...
        inc, ok = _income(vals.get("annual_family_income", ""))
        if not ok:
            errors.append(f"annual_family_income '{vals.get('annual_family_income')}' is not a number")
        out["annual_family_income"] = inc or ""
        cat = vals.get("category", "")
        if cat:
            canon = CATEGORY_CANON.get(_canon(cat))
            if canon:
                out["category"] = canon
            else:
                warns.append(f"category '{cat}' is not one of SC/ST/OBC/General/Minority (kept as is; category-"
                             "restricted schemes will treat it as not matching)")
    return out, errors, warns


def _records(spec: Spec, rows: Iterator[list[str]], rep: Report):
    """Yields (line_no, vals) for data rows; fills header info in rep."""
    header = None
    line = 0
    for row in rows:
        line += 1
        if header is None:
            if not any(row):
                continue
            header = row
            idx = _map_header(spec, header, rep)
            if rep.missing_columns or rep.duplicate_columns:
                return
            continue
        if not any(row):
            continue
        vals = {c: (row[i] if i < len(row) else "") for c, i in idx.items()}
        for c in spec.optional:
            vals.setdefault(c, "")
        extra_vals = len(row) > len(header) and any(row[len(header):])
        yield line, vals, (f"row has {len(row)} values but the header has {len(header)} columns"
                           if extra_vals else None)
    if header is None:
        raise UploadError("the file has no header row")


def validate(spec: Spec, fileobj: IO[bytes], filename: str, max_rows: int = 1_000_000) -> Report:
    rep = Report(spec.dataset, spec.label, filename)
    rep.file_format = detect_format(fileobj, filename)
    encodings = ["-"] if rep.file_format == "xlsx" else ["utf-8-sig", "cp1252"]
    for enc in encodings:
        try:
            _validate_pass(spec, fileobj, rep, enc, max_rows)
            rep.encoding = "" if enc == "-" else ("utf-8" if enc == "utf-8-sig" else enc)
            break
        except UnicodeDecodeError:
            fresh = Report(spec.dataset, spec.label, filename, file_format=rep.file_format)
            rep.__dict__.update(fresh.__dict__)
    else:
        raise UploadError("could not read the text encoding. In Excel use Save As > 'CSV UTF-8 (Comma delimited)'")
    if rep.encoding == "cp1252":
        rep.warnings.insert(0, "file was not UTF-8; read as Windows-1252. Hindi text would be lost: save as "
                               "'CSV UTF-8' if names contain Hindi.")
    if rep.extra_columns:
        rep.warnings.insert(0, f"extra column(s) ignored: {', '.join(rep.extra_columns)}")
    return rep


def _validate_pass(spec, fileobj, rep, enc, max_rows):
    today = date.today()
    keys: set = set()
    warn_counts: dict = {}
    for line, vals, shape_err in _records(spec, iter_rows(fileobj, rep.file_format, enc), rep):
        rep.rows_read += 1
        if rep.rows_read > max_rows:
            rep.too_many_rows = max_rows
            rep.rows_read = max_rows
            break
        _, errors, warns = clean_row(spec, vals, today)
        if shape_err:
            errors.append(shape_err)
        k = vals.get(spec.key, "")
        if k:
            if k in keys:
                errors.append(f"{spec.key} '{k}' appears more than once")
            else:
                keys.add(k)
        for w in warns:
            wk = re.sub(r"'[^']*'", "'…'", w)
            n, ex = warn_counts.get(wk, (0, []))
            val = re.search(r"'([^']*)'", w)
            if val and len(ex) < 5 and val.group(1) not in ex:
                ex.append(val.group(1))
            warn_counts[wk] = (n + 1, ex)
        if errors:
            rep.bad_row_count += 1
            for e in errors:
                p = re.sub(r"'[^']*'", "'…'", e)
                rep.problem_counts[p] = rep.problem_counts.get(p, 0) + 1
            if len(rep.bad_rows_sample) < MAX_BAD_ROW_SAMPLE:
                rep.bad_rows_sample.append({"row": line, "errors": errors,
                                            "values": {c: vals.get(c, "") for c in spec.must_fill}})
    for w, (n, ex) in warn_counts.items():
        rep.warnings.append(f"{n:,} row(s): {w} Values seen: {', '.join(ex)}")


# ------------------------------------------------------------------ loading
COLS = {
    "cbse_results": ("roll_no", "exam_year", "apaar_id", "candidate_name", "dob_raw", "dob_iso", "gender",
                     "father_name", "mother_name", "class_passed", "school", "district", "mobile", "loaded_at"),
    "jan_aadhaar_members": ("member_id", "jan_aadhaar_id", "name_eng", "name_hnd", "dob_raw", "dob_iso", "gender",
                            "gender_norm", "father_name_eng", "mother_name_eng", "category", "annual_family_income",
                            "district", "domicile_state", "disability", "mobile", "relation", "loaded_at"),
}


def _nz(v):
    return v if v not in (None, "") else None


def _db_rows(spec: Spec, fileobj, rep: Report, now) -> Iterator[tuple]:
    today = date.today()
    enc = "utf-8-sig" if rep.encoding in ("", "utf-8") else rep.encoding
    for _, vals, _ in _records(spec, iter_rows(fileobj, rep.file_format, enc), Report(spec.dataset, "", "")):
        v, _, _ = clean_row(spec, vals, today)
        if spec is CBSE:
            yield (v["roll_no"], EXAM_YEAR, _nz(v["apaar_id"]), v["candidate_name"], v["dob"], v["_dob_iso"],
                   v["gender"], _nz(v["father_name"]), _nz(v["mother_name"]), v["class_passed"], _nz(v["school"]),
                   _nz(v["district"]), _nz(v["mobile"]), now)
        else:
            yield (v["member_id"], v["jan_aadhaar_id"], v["nameEng"], _nz(v["nameHnd"]), v["dob"], v["_dob_iso"],
                   v["gender"], v["_gender_norm"], _nz(v["fatherNameEng"]), _nz(v["motherNameEng"]),
                   _nz(v["category"]), _nz(v["annual_family_income"]), _nz(v["district"]), _nz(v["domicile_state"]),
                   _nz(v["disability"]), _nz(v["mobile"]), _nz(v["relation"]), now)


def set_source(conn, dataset: str, source: str, filename, row_count, when, checksum=None, note=None):
    conn.execute(text("DELETE FROM dataset_sources WHERE dataset = :d"), {"d": dataset})
    conn.execute(text("INSERT INTO dataset_sources (dataset, source, filename, row_count, uploaded_at, checksum, note) "
                      "VALUES (:d, :s, :f, :n, :t, :c, :o)"),
                 {"d": dataset, "s": source, "f": filename, "n": row_count, "t": when, "c": checksum, "o": note})


def replace_table(engine: Engine, spec: Spec, fileobj, rep: Report) -> dict:
    """Replace spec.dataset with the (already validated) file in ONE transaction. Also clears the ground
    truth (it cannot describe user data) and old pipeline results (they belong to the old data)."""
    assert rep.ok, "validate() first"
    checksum = sha256_of(fileobj)
    now = datetime.now(timezone.utc)
    with engine.begin() as conn:
        busy = conn.execute(text("SELECT run_id FROM pipeline_runs WHERE status IN ('QUEUED','RUNNING')")).first()
        if busy:
            raise UploadBusy(f"pipeline run {busy[0]} is in progress; wait until it finishes, then upload again")
        gt_before = conn.execute(text("SELECT COUNT(*) FROM ground_truth")).scalar()
        for t in RESULT_TABLES:
            conn.execute(text(f"DELETE FROM {t}"))
        superseded = conn.execute(text("UPDATE pipeline_runs SET status='SUPERSEDED', "
                                       "error='results removed: input data replaced by an upload' "
                                       "WHERE status='SUCCEEDED'")).rowcount
        conn.execute(text("DELETE FROM ground_truth"))
        conn.execute(text(f"DELETE FROM {spec.dataset}"))
        n = bulk_insert(conn, spec.dataset, COLS[spec.dataset], _db_rows(spec, fileobj, rep, now))
        if n != rep.rows_read:
            raise RuntimeError(f"loaded {n} rows but validated {rep.rows_read}; nothing was changed")
        conn.execute(text("DELETE FROM seed_meta WHERE dataset IN (:d, 'ground_truth')"), {"d": spec.dataset})
        set_source(conn, spec.dataset, "uploaded", rep.filename, n, now, checksum)
        set_source(conn, "ground_truth", "cleared", None, 0, now, None,
                   f"cleared by the {spec.label} upload (real data has no ground truth)")
    return {"rows_loaded": n, "uploaded_at": now.isoformat(), "checksum": checksum,
            "ground_truth_rows_cleared": gt_before, "old_runs_superseded": superseded}


# ------------------------------------------------------------------ status
STATUS_TABLES = ("cbse_results", "jan_aadhaar_members", "ground_truth")


def ground_truth_available(conn) -> bool:
    n = conn.execute(text("SELECT COUNT(*) FROM cbse_results")).scalar()
    if not n:
        return False
    uncovered = conn.execute(text("SELECT COUNT(*) FROM cbse_results c LEFT JOIN ground_truth g "
                                  "ON g.roll_no = c.roll_no WHERE g.roll_no IS NULL")).scalar()
    return uncovered == 0


def status(conn) -> dict:
    src = {r["dataset"]: dict(r) for r in conn.execute(text("SELECT * FROM dataset_sources")).mappings()}
    out = {}
    for t in STATUS_TABLES:
        rows = conn.execute(text(f"SELECT COUNT(*) FROM {t}")).scalar()
        s = src.get(t) or {}
        source = s.get("source") or ("empty" if rows == 0 else "unknown")
        if rows == 0 and source == "synthetic":
            source = "empty"
        out[t] = {"rows": rows, "source": source, "filename": s.get("filename"),
                  "rows_at_load": s.get("row_count"), "loaded_at": s.get("uploaded_at"),
                  "checksum": s.get("checksum"), "note": s.get("note")}
    return out
