"""Excel workbook writer (colour-coded, frozen headers). Uses openpyxl write-only mode for speed."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config

GREEN, LGREEN, YELLOW, ORANGE, RED, GREY, BLUE = ("C6EFCE", "E2F0D9", "FFEB9C", "FCD5B4", "FFC7CE",
                                                 "E7E6E6", "DDEBF7")
HEADER_FILLS = {"Read Me": "404040", "Funnel": "1F4E78", "Linkage Decisions": "2F5597",
                "Match Quality": "548235", "Student Eligibility Summary": "7030A0",
                "Eligibility Audit": "7030A0", "Compiled Scheme Rules": "C55A11",
                "Outreach Queue": "BF8F00", "Scenario Coverage": "2F5597", "Scenario Outcomes": "2F5597",
                "Match Outcomes": "548235", "36-Scenario Matrix": "808080", "Top Eligible Schemes": "7030A0"}

ACTION_RULES = [("Auto-link", GREEN), ("Auto-link+flag", YELLOW),
                ("Do not link + Discovery outreach", ORANGE), ("Do not link", RED)]


def _fill(c):
    return PatternFill("solid", fgColor=c)


def _sheet(wb, name, rows, columns=None, widths=None, rules=None, wrap_cols=(), note=None):
    ws = wb.create_sheet(name)
    color = HEADER_FILLS.get(name, "1F4E78")
    ws.sheet_properties.tabColor = color
    columns = columns or (list(rows[0].keys()) if rows else ["(no rows)"])
    widths = widths or {}
    for i, col in enumerate(columns, start=1):
        default = min(max(len(str(col)) + 2, 10), 45)
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, default)
    start_row = 1
    if note:
        ws.freeze_panes = "A3"
        c = WriteOnlyCell(ws, value=note)
        c.font = Font(italic=True, color="C00000")
        ws.append([c])
        start_row = 2
    else:
        ws.freeze_panes = "A2"
    header = []
    for col in columns:
        c = WriteOnlyCell(ws, value=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _fill(color)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        header.append(c)
    ws.append(header)
    n = len(rows)
    first, last = start_row + 1, start_row + max(n, 1)
    # conditional formatting (cheap even for many rows)
    for col, pairs in (rules or {}).items():
        if col not in columns:
            continue
        L = get_column_letter(columns.index(col) + 1)
        rng = f"{L}{first}:{L}{last}"
        for value, fillc in pairs:
            if value.startswith("="):
                formula = value[1:].replace("{c}", f"${L}{first}")
            else:
                formula = f'${L}{first}="{value}"'
            ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=_fill(fillc), stopIfTrue=True))
    wrap_idx = {columns.index(c) for c in wrap_cols if c in columns}
    for r in rows:
        vals = []
        for i, col in enumerate(columns):
            v = r.get(col)
            if isinstance(v, (list, tuple, set, frozenset)):
                v = "; ".join(map(str, v))
            if i in wrap_idx:
                c = WriteOnlyCell(ws, value=v)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                v = c
            vals.append(v)
        ws.append(vals)
    return ws


def write_workbook(path, funnel, decisions, quality, summaries, audits, rules, queue, coverage,
                   top_schemes, results, scenario_outcomes=None,
                   match_outcomes=None, scenarios36=None):
    wb = Workbook(write_only=True)
    st = results["eligibility_stats"]
    rs = results["rule_stats"]
    cap = config.AUDIT_XLSX_ROW_CAP
    readme = [
        ("Product", "MoSJE Product 1 - Scholarship Intelligence & Outreach Platform (prototype, SYNTHETIC data)"),
        ("Linkage rules", "MoSJE Record Linkage Rules V3.0 (implemented as written)"),
        ("Eligibility rules", "MoSJE Scholarship Eligibility Rule V3.0 (implemented as written)"),
        ("Rule_Version", config.RULE_VERSION),
        ("Eligibility_As_Of_Date", results["eligibility_as_of_date"]),
        ("Generated at", results["generated_at"]),
        ("Scheme master", f"{rs['total_rows']} rows; {rs['excluded']} excluded as SUPERSEDED / AGGREGATE "
                          f"PLACEHOLDER (data hygiene, not a rule change); {rs['active']} active"),
        ("Candidate schemes / student", f"{st['candidate_schemes_per_student']} (Central + Rajasthan, active)"),
        ("Eligibility Audit sheet", f"Capped at the first {cap:,} of {st['audit_pairs']:,} student-scheme rows. "
                                    "Full audit: output/eligibility_audit_full.csv"),
        ("Eligible template", "The Linkage doc contains only the Discovery template. The eligible-shortlist "
                              "message is a PROTOTYPE DRAFT written for this demo and needs approval."),
        ("Colour key", "Green = Auto-link / PASS / ELIGIBLE; Yellow = Auto-link+flag; Orange = Discovery "
                       "(probable); Red = Do not link / FAIL / UNRESOLVABLE; Grey = not evaluated / excluded"),
        ("Scenario Outcomes sheet", "All 56 matrix rows + exact-match path + 'no blocking candidate' + default "
                                    "rule, with record counts, TP/FP/FN/TN vs ground truth and G7-blocked counts. "
                                    "Unreachable rows are labelled 'Unreachable'."),
        ("Match Outcomes sheet", "The six headline linkage outcomes vs ground truth (true matches found, false "
                                 "matches, missed matches, precision, recall, false-positive rate), each with its "
                                 "formula and a plain-English definition."),
        ("36-Scenario Matrix sheet", "REPORTING VIEW ONLY. The earlier rule engine's 36-row matrix (verbatim), with "
                                     "every record's V3.0 field scores classified against it (first matching row "
                                     "wins) and ground-truth outcomes of the link V3.0 actually made. V3.0 decides "
                                     "all links; the 36-row action is shown for comparison only."),
        ("Data", "All people, IDs and mobiles are fictional (seeded generator, seed %d)." % config.SEED),
    ]
    _sheet(wb, "Read Me", [{"Item": a, "Detail": b} for a, b in readme], widths={"Item": 28, "Detail": 120},
           wrap_cols=("Detail",))
    _sheet(wb, "Funnel", funnel, widths={"Stage": 48, "Count": 12, "Note": 70})

    dec_rules = {"Final_Action": ACTION_RULES, "Matrix_Action": ACTION_RULES,
                 "GT_Correct": [("YES", GREEN), ("NO", RED)],
                 "G7": [("PASS", GREEN), ("FAIL", RED)],
                 "Linkage_Band": [("MATCHED", GREEN), ("PROBABLE", ORANGE), ("NOT MATCHED", RED)],
                 "DOB": [("Exact", GREEN), ("Mismatch", RED)],
                 "Name_Sim_%": [("=AND(ISNUMBER({c}),{c}>=90)", GREEN), ("=AND(ISNUMBER({c}),{c}>=80)", YELLOW),
                                ("=ISNUMBER({c})", RED)],
                 "Father_Sim_%": [("=AND(ISNUMBER({c}),{c}>=90)", GREEN), ("=AND(ISNUMBER({c}),{c}>=80)", YELLOW),
                                  ("=ISNUMBER({c})", RED)],
                 "Mother_Sim_%": [("=AND(ISNUMBER({c}),{c}>=90)", GREEN), ("=AND(ISNUMBER({c}),{c}>=80)", YELLOW),
                                  ("=ISNUMBER({c})", RED)]}
    _sheet(wb, "Linkage Decisions", decisions, rules=dec_rules,
           widths={"CBSE_Name": 26, "Best_JA_Name": 26, "CBSE_Father": 24, "CBSE_Mother": 22,
                   "Best_JA_Father": 24, "Best_JA_Mother": 22, "Flag_Reason": 60, "Classification": 34,
                   "Final_Action": 30, "Matrix_Action": 30})
    _sheet(wb, "Match Quality", quality, widths={"Segment": 42},
           note="Record-level: TP=linked to true member; FP=linked to wrong member or student not in JA; "
                "FN=true member exists but not linked to it; TN=not in JA and not linked. "
                "Precision=TP/(TP+FP), Recall=TP/(TP+FN), FPR=FP/(FP+TN). Noise segments overlap.")
    if match_outcomes:
        _sheet(wb, "Match Outcomes", match_outcomes,
               columns=["Outcome", "Value", "Formula", "Definition", "Better_When"],
               widths={"Outcome": 24, "Value": 12, "Formula": 18, "Definition": 100, "Better_When": 13},
               wrap_cols=("Definition",),
               note="Overall linkage outcomes vs ground truth (one CBSE record = one decision). TP = linked to the "
                    "true member; FP = linked to the wrong member or student not in Jan Aadhaar; FN = true member "
                    "exists but not linked to it; TN = not in Jan Aadhaar and not linked.")
    sum_cols = ["Student_ID", "Member_ID", "Candidate_Name", "Class_Passed", "Student_Education_Stage",
                "Social_Category", "Gender", "Annual_Family_Income", "Calculated_Age", "Domicile_State_UT",
                "District", "Eligibility_As_Of_Date", "Total_Candidate_Schemes_Checked", "Eligible_Scheme_Count",
                "Eligible_Scheme_IDs", "Eligible_Scheme_Names", "Not_Eligible_Scheme_Count", "Rule_Version",
                "Outreach_Status"]
    _sheet(wb, "Student Eligibility Summary", summaries, columns=sum_cols,
           widths={"Eligible_Scheme_IDs": 50, "Eligible_Scheme_Names": 80, "Candidate_Name": 26},
           rules={"Outreach_Status": [("QUEUE_FOR_OUTREACH", GREEN), ("NO_OUTREACH", RED)]})
    res_rules = [("PASS", GREEN), ("FAIL", RED), ("NOT_EVALUATED (fail-fast)", GREY)]
    audit_cols = ["Student_ID", "Scheme_ID", "Scheme_Name", "Scheme_Level", "Scheme_State_UT",
                  "Domicile_Result", "Education_Stage_Result", "Social_Category_Result", "Gender_Result",
                  "Income_Result", "Age_Result", "Final_Result", "Failure_Reason", "Rule_Version",
                  "Evaluation_Timestamp"]
    _sheet(wb, "Eligibility Audit", audits[:cap], columns=audit_cols,
           widths={"Scheme_Name": 50, "Failure_Reason": 70},
           rules={**{c: res_rules for c in audit_cols if c.endswith("_Result") and c != "Final_Result"},
                  "Final_Result": [("ELIGIBLE", GREEN), ("NOT ELIGIBLE", RED)]},
           note=(f"Showing first {min(cap, len(audits)):,} of {len(audits):,} student x candidate-scheme rows "
                 f"(cap {cap:,}). Full audit in output/eligibility_audit_full.csv. Checks run in bulk order "
                 "(Jurisdiction, Education Stage, Category, Gender, Income, Age) with fail-fast."))
    stat_rules = [("NO_REQUIREMENT", BLUE), ("PARSED", GREEN), ("UNRESOLVABLE", RED)]
    _sheet(wb, "Compiled Scheme Rules", rules,
           widths={"Scheme_Name": 50, "Age_Raw": 40, "Category_Raw": 36, "Category_Note": 50,
                   "Verification_Status": 40, "Exclusion_Reason": 50, "Benefits": 60, "Documents_Required": 60,
                   "Income_Raw": 30},
           rules={"Education_Stage_Status": stat_rules, "Age_Status": stat_rules, "Gender_Status": stat_rules,
                  "Income_Status": stat_rules, "Category_Status": stat_rules,
                  "Active_Status": [("ACTIVE", GREEN), ("EXCLUDED", GREY)],
                  "Compile_Status": [("COMPLETE", GREEN), ("HAS_UNRESOLVABLE", RED), ("EXCLUDED", GREY)]})
    q_cols = ["Outreach_Status", "Message_Type", "Template_Name", "Roll_No", "Candidate_Name", "First_Name",
              "Class_Passed", "Linked_Member_ID", "Mobile", "Contact_Source", "Sendable",
              "Eligible_Scheme_Count", "Top_Schemes_In_Message", "Best_Candidate_Score", "Scenario",
              "Message_EN", "Message_HI", "Eligible_Scheme_IDs"]
    _sheet(wb, "Outreach Queue", queue, columns=q_cols,
           widths={"Message_EN": 70, "Message_HI": 60, "Top_Schemes_In_Message": 60, "Template_Name": 40,
                   "Eligible_Scheme_IDs": 40, "Sendable": 30, "Candidate_Name": 24},
           rules={"Outreach_Status": [("QUEUE_FOR_OUTREACH", GREEN), ("QUEUE_FOR_DISCOVERY_OUTREACH", ORANGE)],
                  "Sendable": [("YES", GREEN), ('=LEFT({c},2)="NO"', RED)]})
    _sheet(wb, "Scenario Coverage", coverage,
           widths={"Classification": 40, "Action": 32, "Achievable_Score_Range": 16},
           rules={"Action": ACTION_RULES, "Mathematically_Reachable": [("NO", GREY), ("Yes", LGREEN)],
                  "Records": [("=AND(ISNUMBER({c}),{c}>0)", YELLOW)]},
           note="Achievable_Score_Range = min/max overall score possible given the row's field bands and "
                "weights (DOB 30, Name 25, Father 20, Mother 20, Gender 5). Rows marked NO can never be hit.")
    if scenario_outcomes:
        from .scenarios import display_rows, totals, COUNT_COLS
        num = [dict(r) for r in scenario_outcomes]
        shown = display_rows(num) + [totals(num)]
        cols = list(scenario_outcomes[0].keys())
        _sheet(wb, "Scenario Outcomes", shown, columns=cols,
               widths={"Criteria": 70, "Classification": 36, "Action": 30, "Path": 30, "Status": 20,
                       "Achievable_Score_Range": 16},
               rules={"Action": [("Auto-link", GREEN), ("Auto-link+flag", YELLOW),
                                 ("Do not link + Discovery", ORANGE), ("Do not link", RED)],
                      "Status": [("Unreachable", GREY), ("Hit", LGREEN)],
                      "FP": [("=AND(ISNUMBER({c}),{c}>0)", RED)], "FN": [("=AND(ISNUMBER({c}),{c}>0)", ORANGE)],
                      "G7_Blocked": [("=AND(ISNUMBER({c}),{c}>0)", YELLOW)]},
               note="One row per matrix scenario (1-56) plus the exact-match path, 'no blocking candidate' and the "
                    "default rule. Ground truth: TP = linked to the true member; FP = linked to the wrong person; "
                    "FN = not linked although the student is in Jan Aadhaar; TN = not linked and genuinely not in "
                    "Jan Aadhaar (TP+FP+FN+TN = Records). 'Unreachable' = the row can never be hit with weights "
                    "30/25/20/20/5. Scenario 1 here counts fuzzy-path records only; exact matches are row EXACT.")
    if scenarios36:
        from .scenarios36 import COUNT_COLS as S36_COUNTS, EXPORT_COLS as S36_COLS, UNREACHABLE as S36_UNR
        shown = [dict(r) for r in scenarios36]
        for r in shown:
            if r["Status"] == S36_UNR:
                for c in S36_COUNTS + ["of which exact path"]:
                    r[c] = S36_UNR
        tot = {"#": "TOTAL", **{c: sum(r[c] for r in scenarios36) for c in S36_COUNTS},
               "Action differs from V3.0 (records)": sum(r["Action differs from V3.0 (records)"] for r in scenarios36)}
        _sheet(wb, "36-Scenario Matrix", shown + [tot], columns=S36_COLS,
               widths={"Classification": 40, "Action": 18, "Status": 16, "V3.0 scenario(s) by conditions": 20,
                       "V3.0 scenario(s) observed": 30, "Note": 80, "Candidate Name": 12, "Overall Score": 12,
                       "Father Name": 12, "Mother Name": 12},
               wrap_cols=("Note", "Classification"),
               rules={"Action": [("Auto-link", GREEN), ("Auto-link+flag", YELLOW), ("Do not link", RED)],
                      "Status": [("Unreachable", GREY)],
                      "False matches (wrong link)": [("=AND(ISNUMBER({c}),{c}>0)", RED)],
                      "Missed matches": [("=AND(ISNUMBER({c}),{c}>0)", ORANGE)],
                      "Action differs from V3.0 (records)": [("=AND(ISNUMBER({c}),{c}>0)", YELLOW)]},
               note="REPORTING VIEW ONLY - V3.0 (56-row matrix + G7) decides every link. Columns # to Action are the "
                    "earlier 36-row matrix verbatim. Each record's V3.0 field scores are checked against rows 1-36 "
                    "top to bottom, first match wins. Counts = ground truth vs the link V3.0 actually made. "
                    "'Unreachable' = cannot occur with weights 30/25/20/20/5. Action comparison applies the doc's "
                    "top-candidate margin rule (>=5) to rows 1-5.")
    _sheet(wb, "Top Eligible Schemes", top_schemes, widths={"Scheme_Name": 70})
    wb.save(path)
